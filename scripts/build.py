#!/usr/bin/env python3
"""Genera el sitio Bitácora de Opiniones OECE a partir del Excel de la bitácora.

Uso:
    python scripts/build.py                      # usa el .xlsx más reciente de data/
    python scripts/build.py "data/archivo.xlsx"  # usa un Excel concreto

Salidas:
    data/opiniones.json   datos limpios
    index.html            sitio listo para publicar (GitHub Pages, Netlify, etc.)

Requiere: pip install openpyxl
"""
import base64
import datetime
import glob
import json
import os
import re
import sys
import unicodedata
from collections import Counter

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LINKEDIN = "https://www.linkedin.com/in/j-rodolfo-mercado-pajares-aab7101ba/"
SHEET = "Bitacora"
FIRST_ROW = 5  # la fila 4 tiene los encabezados
# Página de la opinión en gob.pe: el número inicial del nombre del PDF es el id de la publicación
FICHA = "https://www.gob.pe/institucion/oece/informes-publicaciones/{id}-opinion-n-{code}-{year}-oece-dtn"
GOBPE = os.path.join(ROOT, "data", "gobpe.json")                # opiniones descargadas de gob.pe (scripts/gobpe.py)
CLASIF = os.path.join(ROOT, "data", "clasificacion.json")       # su clasificación por etapa, categoría y tema
BASE_XLSX = os.path.join(ROOT, "base_opiniones.xlsx")           # la base completa, para abrirla en Excel

# Categorías en el orden del ciclo de contratación (hoja "Apoyo" del Excel) -> etiqueta legible
CAT = {
    "Estudios de Pre Inversión": "Estudios de preinversión",
    "Exped._Técnico": "Expediente técnico",
    "Actos_Prepar.": "Actos preparatorios",
    "Proced_Selección": "Procedimiento de selección",
    "Supervis._Inspect.": "Supervisión / inspección",
    "Solo construcción": "Solo construcción",
    "Diseño Costrucción": "Diseño y construcción",
    "Solo Constr/Diseño Constr.": "Construcción / diseño y construcción",
    "Operac_Mantenim.": "Operación y mantenimiento",
    "Compras Corporativas": "Compras corporativas",
    "Bienes y Servicios": "Bienes y servicios",
    "Actos Admin_Entidad": "Actos administrativos de la entidad",
}

# Temas -> etiqueta legible
TEMA = {
    "Supervisor - Inspector": "Supervisor / inspector", "Impedimentos": "Impedimentos",
    "Evaluac.Propuest.": "Evaluación de propuestas", "Adicionales-Deductivos": "Adicionales y deductivos",
    "Controversias": "Controversias", "Otros Procedim_selección": "Otros · procedimiento de selección",
    "Adelantos": "Adelantos", "Penalidades/Multas": "Penalidades y multas", "Mayores Metrados": "Mayores metrados",
    "Liquidaciones": "Liquidaciones", "Comité de selección": "Comité de selección",
    "Ampliación de Plazo": "Ampliación de plazo", "Ampliación Excepcional de Plazo": "Ampliación excepcional de plazo",
    "Present.-Propuest.": "Presentación de propuestas", "Otros obras": "Otros · obras",
    "Cronogr_Programa de Obra": "Cronograma / programa de obra", "Nulidad Contractual": "Nulidad contractual",
    "Garantías": "Garantías", "Actores en el P_S": "Actores del procedimiento de selección",
    "Sustitución de personal técnico": "Sustitución de personal técnico", "Recepción de Obras": "Recepción de obras",
    "Residente de Obra": "Residente de obra", "Inicio de obra": "Inicio de obra",
    "Valorizac_metrados": "Valorizaciones y metrados", "Formul_Polinóm, y Reajustes": "Fórmulas polinómicas y reajustes",
    "Resolución Contractual": "Resolución contractual",
    "Consultas en cuaderno de Incidencias": "Consultas en cuaderno de incidencias",
    "Cuaderno de Incidencias": "Cuaderno de incidencias", "Consult_al Exp_Técnico": "Consultas al expediente técnico",
    "Deducrivos/reducciones": "Deductivos / reducciones", "Intervención Económica": "Intervención económica",
    "Modificaciones Contractuales": "Modificaciones contractuales", "Obligaciones Esenciales": "Obligaciones esenciales",
    "Paralizaciones": "Paralizaciones", "Reactivación de Obras": "Reactivación de obras", "Retrasos": "Retrasos",
    "JRD _JPRD": "JRD / JPRD", "Contratos Menores": "Contratos menores", "Fast Track": "Fast track",
    "Saldos de Obras": "Saldos de obra", "Emergencia": "Emergencia", "Cancelación": "Cancelación",
    "No competitivo": "Procedimiento no competitivo", "Gastos Generales": "Gastos generales",
    "Pago anticipado": "Pago anticipado", "Compra Corporativa": "Compra corporativa",
    "Suspensiones de plazo contractual": "Suspensión del plazo contractual",
    "Anuncio de Contratac_Futuro": "Anuncio de contratación futura", "Contratación Directa": "Contratación directa",
    "Ejecución contractual bienes y servicios": "Ejecución contractual de bienes y servicios",
    "Actos preparatorios": "Actos preparatorios", "Convocatoria": "Convocatoria", "Conformidad": "Conformidad",
    "Declarat_Desierto": "Declaratoria de desierto", "Declarat_Nulidad": "Declaratoria de nulidad",
    "Elevac_OSCE": "Elevación al OECE", "Evaluac.Propuest": "Evaluación de propuestas",
    "Fiscalización Posterior": "Fiscalización posterior", "Formulación de consultas": "Formulación de consultas",
    "Formulación de observac.": "Formulación de observaciones", "Integrac._Bases": "Integración de bases",
    "Otorgam-_Buena Pro": "Otorgamiento de la buena pro", "Registro de participantes": "Registro de participantes",
    "Perfeccionamiento de contrato": "Perfeccionamiento del contrato", "Solvencia económica": "Solvencia económica",
}

# Temas del grupo "Ejecución de obras" (primera lista de la hoja "Apoyo"); el resto va a "Selección, bienes y servicios"
OBRAS = {
    "Adicionales-Deductivos", "Ampliación de Plazo", "Ampliación Excepcional de Plazo", "Adelantos",
    "Cronogr_Programa de Obra", "Consult_al Exp_Técnico", "Consultas en cuaderno de Incidencias", "Controversias",
    "Cuaderno de Incidencias", "Deducrivos/reducciones", "Fast Track", "Formul_Polinóm, y Reajustes", "Garantías",
    "Gastos Generales", "Inicio de obra", "Intervención Económica", "JRD _JPRD", "Liquidaciones", "Mayores Metrados",
    "Modificaciones Contractuales", "Nulidad Contractual", "Obligaciones Esenciales", "Paralizaciones",
    "Penalidades/Multas", "Reactivación de Obras", "Recepción de Obras", "Residente de Obra", "Resolución Contractual",
    "Retrasos", "Saldos de Obras", "Suspensiones de plazo contractual", "Sustitución de personal técnico",
    "Supervisor - Inspector", "Valorizac_metrados", "Otros obras",
}

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
         "setiembre": 9, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

warnings = []


def pretty(raw):
    """Etiqueta legible para una categoría/tema nuevo que no está en los diccionarios."""
    s = re.sub(r"[_]+", " ", raw).replace(".", ". ").strip()
    s = re.sub(r"\s+", " ", s)
    return s[:1].upper() + s[1:].lower()


def marco(text):
    if "32069" in text:
        return "32069"
    if "30225" in text:
        return "30225"
    if "1017" in text:
        return "1017"
    warnings.append(f"Marco normativo no reconocido: {text!r} (se usa 32069)")
    return "32069"


def updated_date(ws):
    """Lee 'Actualizad0: 15 de setiembre del 2026' de la cabecera del Excel."""
    for row in ws.iter_rows(min_row=1, max_row=4):
        for c in row:
            if isinstance(c.value, str):
                m = re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+del?\s+(\d{4})", c.value, re.I)
                if m and m.group(2).lower() in MESES:
                    return datetime.date(int(m.group(3)), MESES[m.group(2).lower()], int(m.group(1))).isoformat()
    warnings.append("No se encontró la fecha de actualización en la cabecera; se usa la fecha de hoy")
    return datetime.date.today().isoformat()


def build(xlsx):
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[SHEET]
    rows, temas_usados, cats_usadas, normas = [], {}, {}, {}
    for r in range(FIRST_ROW, ws.max_row + 1):
        g = lambda col: ws[f"{col}{r}"].value  # noqa: E731
        cat, tema, caso, marco_txt, op, link = g("C"), g("D"), g("E"), g("F"), g("G"), g("H")
        if not (cat and tema and caso and op and link):
            # Las filas de relleno ("…..") y las notas del pie no tienen tema, casuística, opinión ni enlace
            if tema or caso or op or link:
                falta = [n for n, v in (("categoría", cat), ("tema", tema), ("casuística", caso),
                                        ("opinión", op), ("enlace", link)) if not v]
                warnings.append(f"Fila {r}: incompleta (falta {', '.join(falta)}); no se incluye")
            continue
        cat, tema, op = str(cat).strip(), str(tema).strip(), str(op).strip().upper()

        # Texto de la consulta: solo se normalizan espacios; el sufijo "- C2, C3" pasa a lista de conclusiones
        q = re.sub(r"\s+", " ", str(caso)).strip()
        concl = []
        m = re.search(r"[\s\-_–,]*(?<![A-Za-z0-9])(C\d+(?:\s*(?:,|y|-{1,2})\s*C?\d+)*)\s*$", q)
        if m:
            concl = [int(x) for x in re.findall(r"\d+", m.group(1))]
            q = q[: m.start()].strip()

        # Número, año y fecha de carga en gob.pe a partir del enlace del PDF
        mm = re.search(r"opinion-([a-z]?\d+)-(\d{4})-oece-dtn\.pdf\?v=(\d+)", str(link), re.I)
        if mm:
            code, year = mm.group(1).upper(), mm.group(2)
            fecha = datetime.datetime.fromtimestamp(int(mm.group(3)), datetime.timezone(datetime.timedelta(hours=-5))).date().isoformat()
        else:
            code, year, fecha = op, str(datetime.date.today().year), updated_date(ws)
            warnings.append(f"Fila {r}: enlace sin patrón conocido ({link}); fecha = actualización")
        pid = re.search(r"/(\d+)-opinion-[a-z]?\d+-\d{4}-oece-dtn\.pdf", str(link), re.I)
        cm = re.fullmatch(r"([A-Z]?)(\d+)", code)
        ficha = FICHA.format(id=pid.group(1), code=f"{cm.group(1).lower()}{int(cm.group(2)):06d}", year=year) \
            if mm and pid and cm else None

        # Enlaces a la ley y al reglamento (columnas I y J); el Excel los actualiza cada quincena
        m = marco(str(marco_txt or ""))
        ley, reg = g("I"), g("J")
        if ley or reg:
            normas.setdefault(m, Counter())[(str(ley or "").strip(), str(reg or "").strip())] += 1

        if cat not in CAT:
            warnings.append(f"Fila {r}: categoría nueva {cat!r}; agréguela a CAT en build.py")
            CAT[cat] = pretty(cat)
        if tema not in TEMA:
            warnings.append(f"Fila {r}: tema nuevo {tema!r}; agréguelo a TEMA (y a OBRAS si es de obras)")
            TEMA[tema] = pretty(tema)
        temas_usados[tema] = TEMA[tema]
        cats_usadas[cat] = CAT[cat]

        rows.append({
            "n": r - FIRST_ROW + 1, "op": op, "y": year, "u": str(link).strip(), "h": ficha,
            "lk": code if code != op else None,  # el enlace abre otra opinión -> aviso ámbar
            "f": fecha, "c": cat, "t": tema, "g": "obras" if tema in OBRAS else "seleccion",
            "q": q, "k": concl, "m": m,
        })

    links = {}
    for m, cnt in normas.items():
        (ley, reg), _ = cnt.most_common(1)[0]
        if len(cnt) > 1:
            warnings.append(f"Marco {m}: hay {len(cnt)} pares distintos de enlaces a ley/reglamento; se usa el más frecuente")
        links[m] = {"ley": ley or None, "reg": reg or None}

    # ---------- Opiniones descargadas de gob.pe (scripts/gobpe.py) ----------
    # El Excel manda en las opiniones que registra; gob.pe aporta los años anteriores, las opiniones que el Excel
    # aún no registra y, para todas, la fecha de la opinión (la del Excel es la de carga del PDF) y su asunto.
    gob = json.load(open(GOBPE, encoding="utf-8")) if os.path.exists(GOBPE) else {"actualizado": None, "opiniones": {}}
    clas = json.load(open(CLASIF, encoding="utf-8")) if os.path.exists(CLASIF) else {}
    ops = gob["opiniones"]

    # Filas cuyo enlace abre otra opinión: se resuelven contra gob.pe y solo queda el aviso si no está claro
    limpias = [(i, int(r["op"][1:])) for i, r in enumerate(rows) if not r["lk"] and r["op"][1:].isdigit()]
    for i, row in enumerate(rows):
        if not row["lk"]:
            continue
        # vecinas: los números de las filas seguras anterior y siguiente (el Excel va en orden)
        antes = max((n for j, n in limpias if j < i), default=None)
        despues = min((n for j, n in limpias if j > i), default=None)
        fallo = resolver_cruce(row, ops, (antes, despues))
        if not fallo:
            warnings.append(f"Fila rotulada {row['op']}: su enlace abre la {row['lk']} y no se pudo determinar "
                            f"cuál es la correcta; se muestra con el aviso ámbar")
            continue
        quien, g = fallo
        if quien == "enlace":  # el rótulo estaba mal: la fila es de la opinión que abre el enlace
            warnings.append(f"Fila rotulada {row['op']}: por su contenido"
                            f"{' y sus conclusiones ' + ', '.join(map(str, row['k'])) if row['k'] else ''}"
                            f" corresponde a la {row['lk']}; se reasignó")
            row["op"] = row["lk"]
        else:  # el rótulo estaba bien: lo que estaba mal era el enlace
            warnings.append(f"Fila {row['op']}: el enlace del Excel abría la {row['lk']}; "
                            f"se usó el PDF de la {row['op']} publicado en gob.pe")
        row.update(u=g["pdf"], h=g["url"], f=g.get("fecha") or row["f"], lk=None)

    for row in rows:
        row.update(s="OECE", o=f"{row['op'][0]}{int(row['op'][1:]):06d}-{row['y']}-OECE-DTN", src="x")
    en_excel = {f"{r['y']}-OECE-{r['op']}" for r in rows}
    for row in rows:
        g = ops.get(f"{row['y']}-OECE-{row['op']}")
        if g and g.get("fecha"):
            row["f"] = g["fecha"]
    sin_clasificar, n = [], len(rows)
    for key in sorted(ops, key=lambda k: (ops[k]["anio"], ops[k]["serie"] != "OSCE", ops[k]["num"])):
        g = ops[key]
        if key in en_excel:
            continue
        if key not in clas:
            sin_clasificar.append(key)
            continue
        qs = g["consultas"] or [{"q": g["asunto"] or g["titulo"], "m": None}]
        cl = clas[key]
        if isinstance(cl, dict):  # consultas ilegibles en el PDF: una sola ficha con el asunto de la opinión
            qs, cl = [{"q": g["asunto"] or g["titulo"], "m": None}], [cl["asunto"]]
        if len(qs) != len(cl):
            warnings.append(f"{key}: {len(qs)} consultas y {len(cl)} clasificaciones; vuelva a clasificarla")
            continue
        for q, c in zip(qs, cl):
            if c.get("omitir"):  # texto que la extracción tomó por consulta y no lo es
                continue
            n += 1
            temas_usados[c["t"]], cats_usadas[c["c"]] = TEMA[c["t"]], CAT[c["c"]]
            rows.append({
                "n": n, "op": g["op"], "y": g["anio"], "s": g["serie"], "o": g["oficial"], "u": g["pdf"], "h": g["url"],
                "lk": None, "f": g["fecha"] or fecha_pdf(g["pdf"]), "c": c["c"], "t": c["t"],
                "g": "obras" if c["t"] in OBRAS else "seleccion", "q": q["q"], "k": [],
                "m": c.get("m") or q.get("m") or g["marco"], "src": "g",
            })
    if sin_clasificar:
        warnings.append(f"{len(sin_clasificar)} opiniones de gob.pe sin clasificar (no se incluyen; ver "
                        f"'python scripts/gobpe.py pendientes'): {', '.join(sin_clasificar[:12])}"
                        f"{' …' if len(sin_clasificar) > 12 else ''}")

    data = {
        "updated": max(d for d in (updated_date(ws), gob["actualizado"]) if d),
        "rows": rows,
        "cats": [[k, v] for k, v in CAT.items() if k in cats_usadas],
        "temas": temas_usados,
        "normas": links,
        "asuntos": {k: g["asunto"] for k, g in ops.items() if g.get("asunto")},
    }
    return data


VACIAS = {"para", "como", "cuando", "donde", "desde", "hasta", "entre", "sobre", "segun", "cual",
          "cuales", "este", "esta", "estos", "estas", "otro", "otra", "otros", "otras", "puede", "pueden",
          "debe", "deben", "sera", "seria", "sus", "mas", "menos", "caso", "casos", "dicho", "dicha"}


def palabras(txt):
    """Raíces de seis letras: así 'observado' y 'observaciones' cuentan como la misma palabra."""
    t = unicodedata.normalize("NFD", txt or "").encode("ascii", "ignore").decode().lower()
    return {w[:6] for w in re.findall(r"[a-z0-9]+", t) if len(w) >= 4 and w not in VACIAS}


def parecido(txt, g):
    """Parecido entre la casuística del Excel y lo publicado en gob.pe (consultas y asunto).

    Se mide con el coeficiente de Dice, que pesa las dos longitudes: contar solo las palabras
    compartidas premiaría a la consulta más larga por el simple hecho de tener más palabras.
    """
    a = palabras(txt)
    if not a:
        return 0.0
    textos = [c["q"] for c in (g.get("consultas") or [])] + [g.get("asunto") or "", g.get("titulo") or ""]
    return max((2 * len(a & palabras(t)) / (len(a) + len(palabras(t))) for t in textos if t), default=0.0)


def resolver_cruce(row, ops, vecinas):
    """El Excel rotula la fila con un número y su enlace abre otra opinión. ¿Cuál de las dos es?

    Tres criterios, en orden de contundencia:
    1. Las conclusiones: una fila que cita la conclusión 7 no puede ser de una opinión que tiene 2.
    2. El parecido del texto con lo publicado, si una candidata gana con holgura.
    3. El orden del Excel, que va por número de opinión: si el número rotulado encaja entre sus
       vecinas y el del enlace no, el rótulo manda y el enlace se copió de otra fila.
    Devuelve (quién acertó, la opinión) o None si no está claro, y entonces sigue el aviso ámbar.
    """
    cands = {}
    for quien, code in (("rotulo", row["op"]), ("enlace", row["lk"])):
        g = ops.get(f"{row['y']}-OECE-{code}")
        if not g:
            return None  # sin la publicación oficial de ambas no hay con qué decidir
        if row["k"] and g.get("conclusiones") and max(row["k"]) > g["conclusiones"]:
            continue  # cita una conclusión que esa opinión no tiene
        cands[quien] = g
    if len(cands) == 1:
        return next(iter(cands.items()))
    if not cands:
        return None
    (q1, p1), (_, p2) = sorted(((q, parecido(row["q"], g)) for q, g in cands.items()), key=lambda x: -x[1])
    if p1 >= .3 and p1 - p2 >= .08:
        return q1, cands[q1]
    antes, despues = vecinas
    encaja = [q for q, g in cands.items()
              if (antes is None or g["num"] >= antes) and (despues is None or g["num"] <= despues)]
    return (encaja[0], cands[encaja[0]]) if len(encaja) == 1 else None


def fecha_pdf(url):
    """Fecha de carga del PDF en gob.pe (parámetro ?v=, en hora de Lima), si la página no trae la de la opinión."""
    m = re.search(r"\?v=(\d+)", url or "")
    tz = datetime.timezone(datetime.timedelta(hours=-5))
    return datetime.datetime.fromtimestamp(int(m.group(1)), tz).date().isoformat() if m else None


def data_uri(nombre):
    """La imagen de CriterIA va incrustada: el sitio sigue siendo un solo archivo."""
    with open(os.path.join(ROOT, "assets", nombre), "rb") as fh:
        return "data:image/webp;base64," + base64.b64encode(fh.read()).decode("ascii")


def exportar_xlsx(data, ruta):
    """La base completa en una hoja de cálculo, para revisarla o filtrarla en Excel o Google Sheets."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consultas"
    marcos = {"32069": "Ley 32069", "30225": "Ley 30225", "1017": "D.L. 1017"}
    ws.append(["Año", "Institución", "Opinión N.°", "Fecha", "Marco normativo", "Etapa", "Categoría", "Tema",
               "Consulta", "Conclusiones", "Asunto de la opinión", "PDF", "Página en gob.pe", "Clasificación"])
    for r in sorted(data["rows"], key=lambda r: (r["y"], r["s"] != "OSCE", int(re.sub(r"\D", "", r["op"])), r["n"])):
        ws.append([int(r["y"]), r["s"], r["o"], r["f"], marcos.get(r["m"], r["m"]),
                   "Ejecución de obras" if r["g"] == "obras" else "Selección, bienes y servicios",
                   CAT.get(r["c"], r["c"]), TEMA.get(r["t"], r["t"]), r["q"], ", ".join(map(str, r["k"])),
                   data["asuntos"].get(f"{r['y']}-{r['s']}-{r['op']}", ""), r["u"], r["h"] or "",
                   "Bitácora (Excel)" if r["src"] == "x" else "Automática"])
    for col, ancho in zip("ABCDEFGHIJKLMN", (6, 11, 26, 11, 13, 24, 26, 30, 80, 12, 50, 40, 40, 16)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


def main():
    if len(sys.argv) > 1:
        xlsx = sys.argv[1]
    else:
        files = sorted(glob.glob(os.path.join(ROOT, "data", "*.xlsx")), key=os.path.getmtime)
        if not files:
            sys.exit("No hay ningún .xlsx en data/. Copie ahí el Excel de la bitácora.")
        xlsx = files[-1]
    data = build(xlsx)

    # newline="\n": el mismo resultado en Windows que en Mac/Linux (sin diferencias falsas en git)
    with open(os.path.join(ROOT, "data", "opiniones.json"), "w", encoding="utf-8", newline="\n") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)

    tpl = open(os.path.join(ROOT, "src", "template.html"), encoding="utf-8").read()
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = (tpl.replace("{{DATA}}", payload).replace("{{LINKEDIN}}", LINKEDIN)
           .replace("{{AVATAR}}", data_uri("criteria-avatar.webp"))
           .replace("{{MASCOTA}}", data_uri("criteria-mascota.webp")))
    with open(os.path.join(ROOT, "index.html"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(html)
    # Versión para el Artifact de claude.ai: el servicio pone su propio doctype, <head> y <body>,
    # así que se publica desde <title> (título, fuentes y estilos) más el contenido del <body>
    head = html[html.index("<title>"):html.index("</head>")]
    body = html[html.index("<body>") + len("<body>"):html.rindex("</body>")]
    with open(os.path.join(ROOT, "artifact.html"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(head + body)
    exportar_xlsx(data, BASE_XLSX)

    ops = {(r["y"], r["s"], r["op"]) for r in data["rows"]}
    por_anio = Counter(y for y, _, _ in ops)
    print(f"Excel: {os.path.basename(xlsx)}")
    print(f"Actualizado al {data['updated']}: {len(data['rows'])} consultas, {len(ops)} opiniones, {len(data['temas'])} temas")
    print("Opiniones por año:", ", ".join(f"{y}: {c}" for y, c in sorted(por_anio.items())))
    mism = [f"{r['op']}→{r['lk']}" for r in data["rows"] if r["lk"]]
    if mism:
        print("Enlaces que abren otra opinión (aviso ámbar):", ", ".join(mism))
    sin_ficha = sorted({r["op"] for r in data["rows"] if not r["h"]})
    if sin_ficha:
        print("Sin enlace a la página de gob.pe (el PDF no sigue el patrón habitual):", ", ".join(sin_ficha))
    for w in warnings:
        print("AVISO:", w)
    print("Listo: index.html, artifact.html, data/opiniones.json y base_opiniones.xlsx")


if __name__ == "__main__":
    main()
